"""Crosswalk rules that decide which PSA row becomes which LGU's label."""
import pandas as pd

from lgu import labels


def test_norm_and_same_name_tolerate_psa_spellings():
    assert labels.norm("City of Ozamiz") == labels.norm("Ozamiz City")
    assert labels._ratio("Possorubio", "Pozorrubio") >= 0.75  # rule matches accept typos at 0.75
    assert labels.same_name("Santa Cruz", "Sta. Cruz")
    assert not labels.same_name("Sampaloc", "San Pablo")


def test_rematch_prefers_province_then_region_and_refuses_ambiguity():
    adm3 = pd.DataFrame({"adm3_pcode": ["PH1606701", "PH1608505", "PH1600305"],
                         "adm3_name": ["Alegria", "Loreto", "Loreto"]})
    # code points at Surigao del Norte (PH16067); the name exists twice elsewhere in the region -> no match
    assert labels._rematch("Loreto", "PH1606718", adm3) is None
    assert labels._rematch("Loreto", "PH1608501", adm3) == "PH1608505"  # same province wins


def test_region_guard_allows_armm_to_barmm_only():
    assert labels._region_ok("153811", "PH1908711")
    assert not labels._region_ok("133906", "PH0405639")


def test_manila_districts_combine_by_population_with_the_upper_bound_se():
    w = {"133901": 3, "133902": 1}
    rows = [labels.SaeRow("133901", "1st district", "Tondo", 2023, 2.0, 10.0, 0.2, None, None, "s"),
            labels.SaeRow("133902", "", "Binondo", 2023, 6.0, None, 0.6, None, None, "s"),
            labels.SaeRow("133901", "1st district", "Tondo", 2021, 2.0, 10.0, 0.2, None, None, "s")]
    (c,) = labels.combine_districts(rows, w)  # 2021 lacks Binondo: no city row
    assert (c.code6, c.year, c.name) == ("133900", 2023, "City of Manila")
    assert c.estimate == 3.0 and abs(c.se - 0.3) < 1e-12 and abs(c.cv - 10.0) < 1e-9
    assert abs(c.ci_lo - (3.0 - labels.Z90 * 0.3)) < 1e-12


# ------------------------------------------------------------------ chaos: workbook shape and values (audit L-C1, L-C9)
import io

import openpyxl
import pytest

YEARS = (2018, 2021, 2023)


def annex(n: int = 1500, insert_population: bool = False, est=lambda i, k: 10.0 + k, header_se="Standard Error"):
    """A synthetic 2023 Annex 1 in PSA's layout (optionally with a population column inserted after the name)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    pre = ["PSGC", "Region/Province", "Municipality/City"] + (["Population"] if insert_population else [])
    groups = ["Poverty Incidence", None, None, "Coefficient of Variation", None, None, header_se, None, None,
              "90% Confidence Interval", None, None, None, None, None]
    ws.append([None, "Annex 1. Statistical Table"])
    ws.append(pre + groups)
    ws.append([])
    ws.append(["ID"] + [None] * (len(pre) - 1) + [None] * 9 + ["Lower Limit", "Upper Limit"] * 3)
    ws.append([None] * len(pre) + list(YEARS) * 3 + [y for y in YEARS for _ in (0, 1)])
    ws.append([None, "National Capital Region (NCR)"])
    for i in range(n):
        vals = [est(i, k) for k in range(3)]
        row = [140000 + i, "Abra" if i == 0 else None, f"Town {i}"] + ([35000 + i] if insert_population else [])
        row += vals + [10.0] * 3 + [v / 10 for v in vals] + [x for v in vals for x in (v - 1, v + 1)]
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue()), read_only=True, data_only=True)


def test_an_inserted_column_is_read_by_header_name_not_position():
    plain = labels.parse_workbook(annex(), "psa-sae-2023")
    shifted = labels.parse_workbook(annex(insert_population=True), "psa-sae-2023")
    assert [(r.year, r.estimate, r.se, r.ci_lo) for r in shifted] == [(r.year, r.estimate, r.se, r.ci_lo) for r in plain]
    assert {r.estimate for r in shifted if r.year == 2018} == {10.0}  # not the population (35,000+)


def test_a_renamed_or_missing_header_is_a_schema_error():
    with pytest.raises(labels.SchemaError, match="se 2018"):
        labels.parse_workbook(annex(header_se="Std. Err."), "psa-sae-2023")


def test_implausible_values_and_truncated_rounds_are_rejected():
    with pytest.raises(labels.SchemaError, match="outside \\[0, 100\\]"):
        labels.parse_workbook(annex(est=lambda i, k: 35000.0 if i == 7 else 10.0), "psa-sae-2023")
    with pytest.raises(labels.SchemaError, match="1499 LGU rows"):
        labels.parse_workbook(annex(n=1499), "psa-sae-2023")


def test_registry_hook_rejects_a_shape_changed_workbook(tmp_path):
    pytest.importorskip("archivekit")
    from archivekit import SchemaError

    from lgu import sources

    wb = openpyxl.Workbook()
    wb.active.append(["PSGC", "Region/Province", "Municipality/City", "Population"])
    wb.active.append([140001, "Abra", "Bangued", 50000])
    f = tmp_path / "candidate"
    wb.save(f)
    with pytest.raises(SchemaError):
        sources.check_sae(f, {"id": "psa-sae-2023"})


def test_psgc_population_uses_the_newest_census_column():
    assert labels.population_column(["PSGC", "2015 Population", "2020 Population"]) == (2, 2020)
    assert labels.population_column(["PSGC", "2020 Population", "2024 Population"]) == (2, 2024)
    with pytest.raises(labels.SchemaError):
        labels.population_column(["PSGC", "Population (2024 POPCEN)"])


def test_crosswalk_gate_flags_new_unmatched_codes_and_lost_labels():
    committed = {"unmatched": [{"code6": "133901"}], "rows_per_round": {"2021": 1598, "2023": 1598}}
    ok = {"unmatched": [{"code6": "133901"}], "rows_per_round": {"2021": 1598, "2023": 1590, "2025": 1600}}
    assert labels.gate_report(ok, committed, {}) == []
    bad = {"unmatched": [{"code6": "133901"}, {"code6": "153801"}], "rows_per_round": {"2021": 1500}}
    problems = labels.gate_report(bad, committed, {})
    assert any("153801" in p for p in problems) and any("2021" in p for p in problems)
    assert any("round 2023: no labels" in p for p in problems)
    assert not any("153801" in p for p in labels.gate_report(bad, committed, {"153801": ""}))
